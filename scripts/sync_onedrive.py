"""Download privado, conversão validada e renovação do token no GitHub Secrets."""
import os, json, gzip, hashlib, datetime, pathlib, base64, sys, unicodedata, re
from urllib.parse import quote
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
def required(name):
    value = os.environ.get(name)
    if not value:
        raise RuntimeError('Configure o secret ' + name)
    return value
def checked(response, label):
    if not response.ok:
        raise RuntimeError(f'{label}: HTTP {response.status_code}. Verifique a conexão e as permissões.')
    return response
def convert(source, modified=None):
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    def number(value):
        return isinstance(value, (int, float)) and not isinstance(value, bool)
    def sheet(name, width, predicate, extra_headers=None):
        if name not in workbook.sheetnames:
            raise RuntimeError('Aba obrigatória ausente: ' + name)
        def header_key(value):
            text = unicodedata.normalize('NFKD', str(value or ''))
            return re.sub(r'[^a-z0-9]', '', ''.join(c for c in text if not unicodedata.combining(c)).lower())
        headers = [header_key(v) for v in next(workbook[name].iter_rows(min_row=1, max_row=1, values_only=True))]
        extra_indices = [next((i for i, key in enumerate(headers) if key in aliases), None) for aliases in (extra_headers or [])]
        rows = []
        for row in workbook[name].iter_rows(min_row=2, values_only=True):
            if not predicate(row):
                continue
            rows.append([v.isoformat() if isinstance(v, (datetime.date, datetime.datetime)) else v for v in list(row[:width]) + [row[i] if i is not None and i < len(row) else None for i in extra_indices]])
        return rows
    items = sheet('dados-brutos', 12, lambda r: number(r[0]), [('responsavelcard', 'responsavel', 'atribuidopara', 'nome'), ('os', 'numeroos', 'idos')])
    hours = sheet('tempogasto', 12, lambda r: number(r[0]))
    if not items or not hours:
        raise RuntimeError('Importação vazia. A base anterior será preservada.')
    metrics = {key: sheet(name, width, lambda r: bool(r[0]) and number(r[2])) for key, name, width in [('cycle','cycletime',4),('lead','leadtime',4),('throughput','vazão',3),('rework','retrabalho',3)]}
    result = {'schemaVersion': 1, 'metadata': {'sourceName': source.name, 'importedAt': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'sourceModifiedAt': modified, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'lastHoursDate': max(str(r[5])[:10] for r in hours if r[5]), 'counts': {'items': len(items), 'hours': len(hours), 'orders': sum(r[5] == 'Ordem de Serviço' for r in items)}}, 'items': items, 'hours': hours, 'aging': sheet('aging',3,lambda r: number(r[1])), **metrics}
    workbook.close()
    return result

def main():
    import requests
    from nacl.public import PublicKey, SealedBox
    client = required('ONEDRIVE_CLIENT_ID')
    refresh = required('ONEDRIVE_REFRESH_TOKEN')
    github = required('AUTOMATION_GITHUB_TOKEN')
    repo = required('GITHUB_REPOSITORY')
    token = checked(requests.post('https://login.microsoftonline.com/consumers/oauth2/v2.0/token', data={'client_id': client, 'grant_type': 'refresh_token', 'refresh_token': refresh, 'scope': 'https://graph.microsoft.com/Files.Read offline_access'}, timeout=60), 'Autorização Microsoft').json()
    # Persistir a renovação criptografada evita depender de um token antigo.
    headers = {'Authorization': 'Bearer ' + github, 'Accept': 'application/vnd.github+json', 'X-GitHub-Api-Version': '2022-11-28'}
    endpoint = 'https://api.github.com/repos/' + repo + '/actions/secrets'
    public = checked(requests.get(endpoint + '/public-key', headers=headers, timeout=60), 'Chave de Secrets').json()
    encrypted = base64.b64encode(SealedBox(PublicKey(base64.b64decode(public['key']))).encrypt(token.get('refresh_token',refresh).encode())).decode()
    checked(requests.put(endpoint + '/ONEDRIVE_REFRESH_TOKEN', headers=headers, json={'encrypted_value': encrypted, 'key_id': public['key_id']}, timeout=60), 'Renovação do secret')
    path = os.environ.get('ONEDRIVE_FILE_PATH') or 'Documents/SES-PE/Metricas/base_metricas.xlsx'
    graph = 'https://graph.microsoft.com/v1.0/me/drive/root:/' + quote(path, safe='/')
    auth = {'Authorization': 'Bearer ' + token['access_token']}
    item = checked(requests.get(graph, headers=auth, timeout=60), 'Localizar planilha').json()
    download = item.get('@microsoft.graph.downloadUrl')
    if not download:
        raise RuntimeError('Microsoft não retornou o endereço de download.')
    source = ROOT / 'work' / 'base_metricas.xlsx'
    source.parent.mkdir(exist_ok=True)
    source.write_bytes(checked(requests.get(download, timeout=120), 'Download da planilha').content)
    result = convert(source, item.get('lastModifiedDateTime'))
    target = ROOT / 'data' / 'snapshot.json.gz'
    # Não criar commit quando o arquivo do OneDrive permanece igual.
    if target.exists():
        previous = json.loads(gzip.decompress(target.read_bytes()))
        if previous['metadata'].get('sha256') == result['metadata']['sha256'] and not (ROOT / 'data' / 'snapshot.json').exists():
            print('Planilha sem alterações; dados preservados.')
            return
    compressed = gzip.compress(json.dumps(result, ensure_ascii=False, separators=(',', ':'), allow_nan=False).encode(), mtime=0)
    temporary = target.with_suffix('.tmp')
    temporary.write_bytes(compressed)
    temporary.replace(target)
    # A saída JSON tem precedência no build: remover cópia antiga ao migrar para gzip.
    (ROOT / 'data' / 'snapshot.json').unlink(missing_ok=True)
    print('Base atualizada:', len(result['items']), 'registros e', len(result['hours']), 'apontamentos.')

if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        # URLs temporárias de download nunca devem aparecer nos logs públicos.
        print(str(error) if isinstance(error, RuntimeError) else 'Falha de conexão ou leitura da planilha. A base anterior foi preservada.', file=sys.stderr)
        sys.exit(1)
